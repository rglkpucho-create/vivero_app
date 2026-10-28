import pytest
import io
import openpyxl
from app.extensions import db
from app.models import Venta, PedidoCompra, Insumo, LoteProduccion, LoteAnimal
from app.services.notificaciones import NotificacionesService
from tests.conftest import login_as


def test_exportacion_excel_ventas_compras_e_inventario(client, app, test_data):
    """
    Test para validar la generación y descarga de los 3 Libros Contables en Excel (.xlsx):
    1. Libro de Ventas
    2. Libro de Compras
    3. Kardex Valorizado de Inventario (con sus 4 hojas de cálculo y balance general)
    """
    login_as(client, 'admin_a@test.com', 'Secret123!')

    # 1. Acceso a la vista principal del Hub de Reportes
    res_index = client.get('/reportes')
    assert res_index.status_code == 200
    assert b"Centro de Reportes" in res_index.data

    # 2. Descarga del Libro de Ventas en Excel
    res_ventas_xl = client.get('/reportes/ventas/excel')
    assert res_ventas_xl.status_code == 200
    assert 'spreadsheetml.sheet' in res_ventas_xl.content_type
    assert len(res_ventas_xl.data) > 2000

    # Validar que openpyxl pueda abrir el archivo generado en memoria
    wb_v = openpyxl.load_workbook(io.BytesIO(res_ventas_xl.data))
    assert "Libro de Ventas" in wb_v.sheetnames
    assert "Resumen Métodos de Pago" in wb_v.sheetnames

    # 3. Descarga del Libro de Compras en Excel
    res_compras_xl = client.get('/reportes/compras/excel')
    assert res_compras_xl.status_code == 200
    assert 'spreadsheetml.sheet' in res_compras_xl.content_type
    assert len(res_compras_xl.data) > 1500

    wb_c = openpyxl.load_workbook(io.BytesIO(res_compras_xl.data))
    assert "Libro de Compras" in wb_c.sheetnames

    # 4. Descarga del Kardex Valorizado de Inventario
    res_kardex_xl = client.get('/reportes/inventario/excel')
    assert res_kardex_xl.status_code == 200
    assert 'spreadsheetml.sheet' in res_kardex_xl.content_type
    assert len(res_kardex_xl.data) > 2500

    wb_k = openpyxl.load_workbook(io.BytesIO(res_kardex_xl.data))
    assert "Balance Consolidado" in wb_k.sheetnames
    assert "Insumos y Mercancía" in wb_k.sheetnames
    assert "Activos Biológicos - Cultivos" in wb_k.sheetnames
    assert "Activos Pecuarios - Animales" in wb_k.sheetnames


def test_notificaciones_whatsapp_y_email_factura(client, app, test_data):
    """
    Test para validar las notificaciones automáticas por WhatsApp y Correo Electrónico:
    1. Generación de enlace enriquecido de WhatsApp para Factura
    2. Generación de enlace de WhatsApp para Pedido Web
    3. Envío de correo electrónico con Factura PDF adjunta (Sandbox / SMTP)
    """
    login_as(client, 'admin_a@test.com', 'Secret123!')

    with app.app_context():
        # Crear venta finalizada de prueba
        v = Venta(
            empresa_id=test_data['empresa_a_id'],
            numero_factura="FAC-NOTIF-TEST-001",
            cliente_id=test_data['cliente_a_id'],
            estado='Finalizada',
            total=85000.0,
            metodo_pago='Efectivo'
        )
        db.session.add(v)
        db.session.commit()
        venta_id = v.id

    # 1. Endpoint para obtener URL de WhatsApp de Factura
    res_wa_fac = client.get(f'/venta/{venta_id}/whatsapp-url')
    assert res_wa_fac.status_code == 200
    data_wa = res_wa_fac.get_json()
    assert data_wa['success'] is True
    assert 'api.whatsapp.com/send' in data_wa['whatsapp_url']
    assert 'FAC-NOTIF-TEST-001' in data_wa['whatsapp_url']

    # 2. Endpoint para obtener URL de WhatsApp de Pedido Web
    res_wa_ped = client.get(f'/pedidos-web/{venta_id}/notificar-whatsapp?estado=Listo+para+Entrega')
    assert res_wa_ped.status_code == 200
    data_ped = res_wa_ped.get_json()
    assert data_ped['success'] is True
    assert 'api.whatsapp.com/send' in data_ped['whatsapp_url']

    # 3. Endpoint para envío de correo con PDF adjunto (JSON)
    res_email = client.post(
        f'/venta/{venta_id}/enviar-email',
        json={
            'email': 'cliente.prueba@vivero.com',
            'mensaje': 'Apreciado cliente, adjuntamos su factura comercial.'
        },
        headers={'X-Requested-With': 'XMLHttpRequest'}
    )
    assert res_email.status_code == 200
    data_email = res_email.get_json()
    assert data_email['success'] is True
    assert 'cliente.prueba@vivero.com' in data_email['message']
