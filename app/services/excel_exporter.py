"""
Servicio de Exportación Contable a Excel (.xlsx) — EcoSistem
Genera reportes financieros, libros diarios de ventas, compras y kardex valorizado de inventario.
"""
import io
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# ==============================================================
# ESTILOS VISUALES CORPORATIVOS (PALETA AGRO / ESMERALDA)
# ==============================================================
COLOR_PRIMARY_DARK = "064E3B"   # Verde bosque oscuro
COLOR_PRIMARY_LIGHT = "ECFDF5"  # Menta suave para subtotales
COLOR_ACCENT = "059669"         # Esmeralda vibrante
COLOR_BORDER = "CBD5E1"         # Gris suave
COLOR_ZEBRA = "F8FAFC"          # Gris ultra claro para filas alternas

FONT_TITLE = Font(name="Calibri", size=15, bold=True, color=COLOR_PRIMARY_DARK)
FONT_SUBTITLE = Font(name="Calibri", size=10, italic=True, color="475569")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_DATA = Font(name="Calibri", size=10, color="0F172A")
FONT_TOTAL = Font(name="Calibri", size=11, bold=True, color=COLOR_PRIMARY_DARK)

FILL_HEADER = PatternFill(start_color=COLOR_PRIMARY_DARK, end_color=COLOR_PRIMARY_DARK, fill_type="solid")
FILL_TOTAL = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color=COLOR_ZEBRA, end_color=COLOR_ZEBRA, fill_type="solid")

BORDER_THIN = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_BORDER),
    bottom=Side(style='thin', color=COLOR_BORDER)
)

BORDER_TOTAL = Border(
    left=Side(style='thin', color=COLOR_BORDER),
    right=Side(style='thin', color=COLOR_BORDER),
    top=Side(style='thin', color=COLOR_PRIMARY_DARK),
    bottom=Side(style='double', color=COLOR_PRIMARY_DARK)
)

ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")
ALIGN_RIGHT = Alignment(horizontal="right", vertical="center")


def _aplicar_autoajuste_columnas(ws, min_width=12, max_width=45):
    """Ajusta automáticamente el ancho de las columnas según su contenido."""
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val_str = str(cell.value or '')
            if cell.number_format and '$' in cell.number_format:
                val_str += '     '
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(min_width, min(max_len + 3, max_width))


class ExcelExporterService:

    @staticmethod
    def generar_libro_ventas(empresa, ventas, fecha_inicio=None, fecha_fin=None):
        """
        Genera el Libro Diario de Ventas y Facturación Fiscal en formato Excel (.xlsx).
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Libro de Ventas"
        ws.views.sheetView[0].showGridLines = True

        # 1. Encabezado de la Empresa
        empresa_nombre = empresa.nombre if empresa else "Vivero El Edén Forestal"
        empresa_nit = empresa.nit_rut if (empresa and empresa.nit_rut) else "901.842.190-4"

        ws.merge_cells("A1:M1")
        ws["A1"] = f"{empresa_nombre.upper()} — NIT: {empresa_nit}"
        ws["A1"].font = FONT_TITLE
        ws["A1"].alignment = ALIGN_LEFT

        ws.merge_cells("A2:M2")
        rango_str = f"Período: {fecha_inicio.strftime('%d/%m/%Y') if fecha_inicio else 'Histórico'} al {fecha_fin.strftime('%d/%m/%Y') if fecha_fin else 'Hoy'}"
        ws["A2"] = f"LIBRO DIARIO DE VENTAS Y FACTURACIÓN FISCAL • {rango_str}"
        ws["A2"].font = FONT_SUBTITLE
        ws["A2"].alignment = ALIGN_LEFT

        # 2. Fila de Encabezados de Tabla
        headers = [
            "Factura #", "Fecha & Hora", "Cliente", "Documento / NIT", "Ciudad",
            "Tipo Venta", "Método de Pago", "Subtotal Base", "Descuento",
            "IVA Agropecuario", "Total Facturado", "Estado DIAN", "CUFE / Transacción"
        ]

        row_header = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_header, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN
        ws.row_dimensions[row_header].height = 26

        # 3. Llenado de Filas de Ventas
        current_row = row_header + 1
        subtotal_sum = 0.0
        descuento_sum = 0.0
        iva_sum = 0.0
        total_sum = 0.0

        resumen_metodos = {}

        for idx, v in enumerate(ventas):
            subtotal_sum += (v.subtotal or 0.0)
            descuento_sum += (v.descuento or 0.0)
            iva_sum += (v.impuestos or 0.0)
            total_sum += (v.total or 0.0)

            # Acumulador por método
            mp = v.metodo_pago or "Efectivo"
            resumen_metodos[mp] = resumen_metodos.get(mp, 0.0) + (v.total or 0.0)

            fill_row = FILL_ZEBRA if idx % 2 == 1 else None

            row_data = [
                v.numero_factura,
                v.fecha.strftime('%Y-%m-%d %H:%M') if v.fecha else '',
                v.cliente.nombre if v.cliente else 'Consumidor Final',
                v.cliente.documento if (v.cliente and v.cliente.documento) else '222222222222',
                v.cliente.ciudad if (v.cliente and v.cliente.ciudad) else '',
                v.tipo_venta or 'Contado',
                v.metodo_pago or 'Efectivo',
                v.subtotal or 0.0,
                v.descuento or 0.0,
                v.impuestos or 0.0,
                v.total or 0.0,
                v.estado_dian or 'NO_EMITIDA',
                v.cufe or v.referencia_pago or ''
            ]

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = FONT_DATA
                cell.border = BORDER_THIN
                if fill_row:
                    cell.fill = fill_row

                # Alineaciones y Formatos
                if col_idx in [1, 2, 4, 6, 12]:
                    cell.alignment = ALIGN_CENTER
                elif col_idx in [8, 9, 10, 11]:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "$#,##0"
                else:
                    cell.alignment = ALIGN_LEFT

            current_row += 1

        # 4. Fila de Totales
        ws.cell(row=current_row, column=1, value="TOTAL CONSOLIDADO").alignment = ALIGN_LEFT
        ws.cell(row=current_row, column=1).font = FONT_TOTAL

        for col_idx in range(1, 14):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = FILL_TOTAL
            cell.border = BORDER_TOTAL
            cell.font = FONT_TOTAL

        ws.cell(row=current_row, column=8, value=subtotal_sum).number_format = "$#,##0"
        ws.cell(row=current_row, column=9, value=descuento_sum).number_format = "$#,##0"
        ws.cell(row=current_row, column=10, value=iva_sum).number_format = "$#,##0"
        ws.cell(row=current_row, column=11, value=total_sum).number_format = "$#,##0"

        _aplicar_autoajuste_columnas(ws)

        # 5. Segunda Hoja: Resumen por Método de Pago
        ws2 = wb.create_sheet(title="Resumen Métodos de Pago")
        ws2.views.sheetView[0].showGridLines = True

        ws2.merge_cells("A1:C1")
        ws2["A1"] = "RESUMEN DE RECAUDO POR MÉTODO DE PAGO"
        ws2["A1"].font = FONT_TITLE

        headers_r = ["Método de Pago", "Total Recaudado ($)", "% Participación"]
        for col_idx, h in enumerate(headers_r, 1):
            cell = ws2.cell(row=3, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN

        row_r = 4
        for metodo, monto in sorted(resumen_metodos.items(), key=lambda x: x[1], reverse=True):
            pct = (monto / total_sum) if total_sum > 0 else 0.0

            c1 = ws2.cell(row=row_r, column=1, value=metodo)
            c2 = ws2.cell(row=row_r, column=2, value=monto)
            c3 = ws2.cell(row=row_r, column=3, value=pct)

            c1.font = FONT_DATA; c1.border = BORDER_THIN
            c2.font = FONT_DATA; c2.border = BORDER_THIN; c2.number_format = "$#,##0"; c2.alignment = ALIGN_RIGHT
            c3.font = FONT_DATA; c3.border = BORDER_THIN; c3.number_format = "0.0%"; c3.alignment = ALIGN_RIGHT
            row_r += 1

        # Total resumen
        c_tot_label = ws2.cell(row=row_r, column=1, value="TOTAL GENERAL")
        c_tot_val = ws2.cell(row=row_r, column=2, value=total_sum)
        c_tot_pct = ws2.cell(row=row_r, column=3, value=1.0 if total_sum > 0 else 0.0)

        for c in [c_tot_label, c_tot_val, c_tot_pct]:
            c.font = FONT_TOTAL; c.fill = FILL_TOTAL; c.border = BORDER_TOTAL
        c_tot_val.number_format = "$#,##0"
        c_tot_pct.number_format = "0.0%"

        _aplicar_autoajuste_columnas(ws2)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generar_libro_compras(empresa, compras, fecha_inicio=None, fecha_fin=None):
        """
        Genera el Libro Diario de Compras y Gastos en formato Excel (.xlsx).
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Libro de Compras"
        ws.views.sheetView[0].showGridLines = True

        empresa_nombre = empresa.nombre if empresa else "Vivero El Edén Forestal"
        empresa_nit = empresa.nit_rut if (empresa and empresa.nit_rut) else "901.842.190-4"

        ws.merge_cells("A1:H1")
        ws["A1"] = f"{empresa_nombre.upper()} — NIT: {empresa_nit}"
        ws["A1"].font = FONT_TITLE

        ws.merge_cells("A2:H2")
        rango_str = f"Período: {fecha_inicio.strftime('%d/%m/%Y') if fecha_inicio else 'Histórico'} al {fecha_fin.strftime('%d/%m/%Y') if fecha_fin else 'Hoy'}"
        ws["A2"] = f"LIBRO DE COMPRAS, GASTOS Y PROVEEDORES • {rango_str}"
        ws["A2"].font = FONT_SUBTITLE

        headers = [
            "Pedido / Factura #", "Fecha Compra", "Proveedor", "NIT / Documento",
            "Sede Destino", "Estado", "Total Compra ($)", "Registrado Por"
        ]

        row_header = 4
        for col_idx, h in enumerate(headers, 1):
            cell = ws.cell(row=row_header, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN
        ws.row_dimensions[row_header].height = 26

        current_row = row_header + 1
        total_compras_sum = 0.0

        for idx, c in enumerate(compras):
            total_val = getattr(c, 'total_estimado', None) or getattr(c, 'total', 0.0) or 0.0
            total_compras_sum += total_val
            fill_row = FILL_ZEBRA if idx % 2 == 1 else None

            f_obj = getattr(c, 'fecha_solicitud', None) or getattr(c, 'fecha', None)
            f_str = f_obj.strftime('%Y-%m-%d') if f_obj else ''

            prov = c.proveedor
            prov_nom = getattr(prov, 'razon_social', '') or getattr(prov, 'nombre', '') if prov else 'Proveedor General'
            prov_doc = getattr(prov, 'nit_rut', '') or getattr(prov, 'documento', '') if prov else ''

            num_ref = getattr(c, 'numero_orden', None) or getattr(c, 'numero_pedido', None) or f"PED-{c.id}"

            row_data = [
                num_ref,
                f_str,
                prov_nom,
                prov_doc,
                c.sede.nombre if c.sede else 'Principal',
                c.estado or 'Completado',
                total_val,
                'Administrador'
            ]

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=current_row, column=col_idx, value=val)
                cell.font = FONT_DATA
                cell.border = BORDER_THIN
                if fill_row: cell.fill = fill_row

                if col_idx in [1, 2, 4, 6]:
                    cell.alignment = ALIGN_CENTER
                elif col_idx == 7:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "$#,##0"
                else:
                    cell.alignment = ALIGN_LEFT

            current_row += 1

        # Total compras
        ws.cell(row=current_row, column=1, value="TOTAL COMPRAS Y GASTOS").alignment = ALIGN_LEFT
        for col_idx in range(1, 9):
            cell = ws.cell(row=current_row, column=col_idx)
            cell.fill = FILL_TOTAL
            cell.border = BORDER_TOTAL
            cell.font = FONT_TOTAL

        ws.cell(row=current_row, column=7, value=total_compras_sum).number_format = "$#,##0"

        _aplicar_autoajuste_columnas(ws)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def generar_kardex_inventario_valorizado(empresa, insumos, lotes_cultivo, lotes_animal):
        """
        Genera el Kardex Valorizado de Inventario (Insumos, Cultivos Biológicos y Animales)
        con hoja ejecutiva de balance para el contador.
        """
        wb = openpyxl.Workbook()

        # ==============================================================
        # HOJA 1: INSUMOS Y MERCANCÍA
        # ==============================================================
        ws_insumos = wb.active
        ws_insumos.title = "Insumos y Mercancía"
        ws_insumos.views.sheetView[0].showGridLines = True

        ws_insumos.merge_cells("A1:J1")
        ws_insumos["A1"] = f"{empresa.nombre.upper() if empresa else 'ECOSISTEM'} — KARDEX VALORIZADO DE MERCANCÍA E INSUMOS"
        ws_insumos["A1"].font = FONT_TITLE

        headers_ins = [
            "SKU / Código", "Nombre del Insumo / Producto", "Categoría", "U. Medida",
            "Stock Actual", "Costo Unit. Promedio", "Valor Total (Costo)",
            "Precio Venta Unit.", "Valor Comercial Total", "Margen Bruto Est."
        ]
        for col_idx, h in enumerate(headers_ins, 1):
            cell = ws_insumos.cell(row=3, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN

        row_ins = 4
        total_costo_insumos = 0.0
        total_comercial_insumos = 0.0

        for idx, item in enumerate(insumos):
            stock = item.cantidad_actual or 0.0
            costo_unit = item.costo_promedio or 0.0
            costo_tot = stock * costo_unit
            precio_vta = item.precio_venta or 0.0
            comercial_tot = stock * precio_vta
            margen_pct = ((precio_vta - costo_unit) / precio_vta) if precio_vta > 0 else 0.0

            total_costo_insumos += costo_tot
            total_comercial_insumos += comercial_tot

            fill_row = FILL_ZEBRA if idx % 2 == 1 else None

            u_medida = 'Unidad'
            if hasattr(item, 'unidad_medida_rel') and item.unidad_medida_rel:
                u_medida = item.unidad_medida_rel.nombre
            elif hasattr(item, 'unidad_medida') and item.unidad_medida:
                u_medida = str(item.unidad_medida)

            vals = [
                item.codigo or f"INS-{item.id}",
                item.nombre,
                getattr(item, 'tipo', 'General') or 'General',
                u_medida,
                stock,
                costo_unit,
                costo_tot,
                precio_vta,
                comercial_tot,
                margen_pct
            ]

            for col_idx, val in enumerate(vals, 1):
                cell = ws_insumos.cell(row=row_ins, column=col_idx, value=val)
                cell.font = FONT_DATA
                cell.border = BORDER_THIN
                if fill_row: cell.fill = fill_row

                if col_idx in [1, 3, 4]:
                    cell.alignment = ALIGN_CENTER
                elif col_idx in [5]:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "#,##0.0"
                elif col_idx in [6, 7, 8, 9]:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "$#,##0"
                elif col_idx == 10:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "0.0%"
                else:
                    cell.alignment = ALIGN_LEFT
            row_ins += 1

        # Totales Insumos
        ws_insumos.cell(row=row_ins, column=2, value="TOTAL MERCANCÍA E INSUMOS").font = FONT_TOTAL
        for c in range(1, 11):
            cell = ws_insumos.cell(row=row_ins, column=c)
            cell.fill = FILL_TOTAL; cell.border = BORDER_TOTAL; cell.font = FONT_TOTAL
        ws_insumos.cell(row=row_ins, column=7, value=total_costo_insumos).number_format = "$#,##0"
        ws_insumos.cell(row=row_ins, column=9, value=total_comercial_insumos).number_format = "$#,##0"
        _aplicar_autoajuste_columnas(ws_insumos)

        # ==============================================================
        # HOJA 2: ACTIVOS BIOLÓGICOS - CULTIVOS
        # ==============================================================
        ws_cultivos = wb.create_sheet(title="Activos Biológicos - Cultivos")
        ws_cultivos.views.sheetView[0].showGridLines = True

        ws_cultivos.merge_cells("A1:H1")
        ws_cultivos["A1"] = f"{empresa.nombre.upper() if empresa else 'ECOSISTEM'} — VALORIZACIÓN DE CULTIVOS Y LOTES"
        ws_cultivos["A1"].font = FONT_TITLE

        headers_cult = [
            "Código Lote", "Especie Común", "Nombre Científico", "Estado",
            "Plantas Actuales", "Costo Acumulado Total", "Precio Venta Sugerido", "Valor Comercial Total"
        ]
        for col_idx, h in enumerate(headers_cult, 1):
            cell = ws_cultivos.cell(row=3, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN

        row_cult = 4
        total_costo_cultivos = 0.0
        total_comercial_cultivos = 0.0

        for idx, lote in enumerate(lotes_cultivo):
            cant = lote.cantidad_actual or 0.0
            costo_acum = lote.costo_total or 0.0
            esp = lote.especie
            precio_sug = esp.precio_venta_sugerido if esp else 0.0
            comercial_tot = cant * precio_sug

            total_costo_cultivos += costo_acum
            total_comercial_cultivos += comercial_tot

            fill_row = FILL_ZEBRA if idx % 2 == 1 else None

            vals = [
                lote.codigo_lote,
                esp.nombre_comun if esp else 'Cultivo Varios',
                esp.nombre_cientifico if esp else '',
                lote.estado or 'En Proceso',
                cant,
                costo_acum,
                precio_sug,
                comercial_tot
            ]

            for col_idx, val in enumerate(vals, 1):
                cell = ws_cultivos.cell(row=row_cult, column=col_idx, value=val)
                cell.font = FONT_DATA
                cell.border = BORDER_THIN
                if fill_row: cell.fill = fill_row

                if col_idx in [1, 4]:
                    cell.alignment = ALIGN_CENTER
                elif col_idx == 5:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "#,##0"
                elif col_idx in [6, 7, 8]:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "$#,##0"
                else:
                    cell.alignment = ALIGN_LEFT
            row_cult += 1

        # Totales Cultivos
        ws_cultivos.cell(row=row_cult, column=2, value="TOTAL CULTIVOS Y ACTIVOS VEGETALES").font = FONT_TOTAL
        for c in range(1, 9):
            cell = ws_cultivos.cell(row=row_cult, column=c)
            cell.fill = FILL_TOTAL; cell.border = BORDER_TOTAL; cell.font = FONT_TOTAL
        ws_cultivos.cell(row=row_cult, column=6, value=total_costo_cultivos).number_format = "$#,##0"
        ws_cultivos.cell(row=row_cult, column=8, value=total_comercial_cultivos).number_format = "$#,##0"
        _aplicar_autoajuste_columnas(ws_cultivos)

        # ==============================================================
        # HOJA 3: ACTIVOS PECUARIOS - ANIMALES
        # ==============================================================
        ws_animales = wb.create_sheet(title="Activos Pecuarios - Animales")
        ws_animales.views.sheetView[0].showGridLines = True

        ws_animales.merge_cells("A1:G1")
        ws_animales["A1"] = f"{empresa.nombre.upper() if empresa else 'ECOSISTEM'} — VALORIZACIÓN DE LOTE PECUARIO / ANIMALES"
        ws_animales["A1"].font = FONT_TITLE

        headers_anim = [
            "Código / Arete", "Nombre / Lote", "Raza / Especie", "Estado",
            "Cabezas / Cantidad", "Costo Acumulado", "Valor Comercial Estimado"
        ]
        for col_idx, h in enumerate(headers_anim, 1):
            cell = ws_animales.cell(row=3, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN

        row_anim = 4
        total_costo_animales = 0.0
        total_comercial_animales = 0.0

        for idx, animal in enumerate(lotes_animal):
            cant = animal.cantidad_actual or 0
            costo = animal.costo_acumulado or 0.0
            raza = animal.raza_rel
            valor_est = (raza.valor_mercado_estimado * cant) if (raza and raza.valor_mercado_estimado) else (costo * 1.25)

            total_costo_animales += costo
            total_comercial_animales += valor_est

            fill_row = FILL_ZEBRA if idx % 2 == 1 else None

            vals = [
                animal.codigo,
                animal.nombre or 'Lote Animal',
                raza.nombre if raza else 'General',
                animal.estado or 'Activo',
                cant,
                costo,
                valor_est
            ]

            for col_idx, val in enumerate(vals, 1):
                cell = ws_animales.cell(row=row_anim, column=col_idx, value=val)
                cell.font = FONT_DATA
                cell.border = BORDER_THIN
                if fill_row: cell.fill = fill_row

                if col_idx in [1, 4]:
                    cell.alignment = ALIGN_CENTER
                elif col_idx == 5:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "#,##0"
                elif col_idx in [6, 7]:
                    cell.alignment = ALIGN_RIGHT
                    cell.number_format = "$#,##0"
                else:
                    cell.alignment = ALIGN_LEFT
            row_anim += 1

        # Totales Animales
        ws_animales.cell(row=row_anim, column=2, value="TOTAL ACTIVOS PECUARIOS").font = FONT_TOTAL
        for c in range(1, 8):
            cell = ws_animales.cell(row=row_anim, column=c)
            cell.fill = FILL_TOTAL; cell.border = BORDER_TOTAL; cell.font = FONT_TOTAL
        ws_animales.cell(row=row_anim, column=6, value=total_costo_animales).number_format = "$#,##0"
        ws_animales.cell(row=row_anim, column=7, value=total_comercial_animales).number_format = "$#,##0"
        _aplicar_autoajuste_columnas(ws_animales)

        # ==============================================================
        # HOJA 4: RESUMEN EJECUTIVO Y BALANCE CONTABLE
        # ==============================================================
        ws_balance = wb.create_sheet(title="Balance Consolidado", index=0)
        ws_balance.views.sheetView[0].showGridLines = True

        ws_balance.merge_cells("A1:D1")
        ws_balance["A1"] = f"{empresa.nombre.upper() if empresa else 'ECOSISTEM'} — BALANCE VALORIZADO DE INVENTARIOS"
        ws_balance["A1"].font = FONT_TITLE

        ws_balance.merge_cells("A2:D2")
        ws_balance["A2"] = f"Generado el {datetime.now().strftime('%d/%m/%Y a las %I:%M %p')} • Moneda: Pesos Colombianos (COP)"
        ws_balance["A2"].font = FONT_SUBTITLE

        headers_bal = ["Tipo de Activo de Inventario", "Valor a Costo Contable ($)", "Valor Comercial Estimado ($)", "Plusvalía / Margen Bruto ($)"]
        for col_idx, h in enumerate(headers_bal, 1):
            cell = ws_balance.cell(row=4, column=col_idx, value=h)
            cell.font = FONT_HEADER
            cell.fill = FILL_HEADER
            cell.alignment = ALIGN_CENTER
            cell.border = BORDER_THIN
        ws_balance.row_dimensions[4].height = 25

        filas_balance = [
            ("1. Mercancía e Insumos Operativos", total_costo_insumos, total_comercial_insumos),
            ("2. Activos Biológicos Vegetales (Cultivos/Plantas)", total_costo_cultivos, total_comercial_cultivos),
            ("3. Activos Biológicos Animales (Pecuario)", total_costo_animales, total_comercial_animales),
        ]

        row_bal = 5
        gran_total_costo = 0.0
        gran_total_comercial = 0.0

        for desc, c_tot, com_tot in filas_balance:
            gran_total_costo += c_tot
            gran_total_comercial += com_tot
            plusvalia = com_tot - c_tot

            c1 = ws_balance.cell(row=row_bal, column=1, value=desc)
            c2 = ws_balance.cell(row=row_bal, column=2, value=c_tot)
            c3 = ws_balance.cell(row=row_bal, column=3, value=com_tot)
            c4 = ws_balance.cell(row=row_bal, column=4, value=plusvalia)

            c1.font = FONT_DATA; c1.border = BORDER_THIN
            c2.font = FONT_DATA; c2.border = BORDER_THIN; c2.number_format = "$#,##0"; c2.alignment = ALIGN_RIGHT
            c3.font = FONT_DATA; c3.border = BORDER_THIN; c3.number_format = "$#,##0"; c3.alignment = ALIGN_RIGHT
            c4.font = FONT_DATA; c4.border = BORDER_THIN; c4.number_format = "$#,##0"; c4.alignment = ALIGN_RIGHT
            row_bal += 1

        # Gran Total
        c_fin_1 = ws_balance.cell(row=row_bal, column=1, value="GRAN TOTAL VALORIZACIÓN DE ACTIVOS")
        c_fin_2 = ws_balance.cell(row=row_bal, column=2, value=gran_total_costo)
        c_fin_3 = ws_balance.cell(row=row_bal, column=3, value=gran_total_comercial)
        c_fin_4 = ws_balance.cell(row=row_bal, column=4, value=gran_total_comercial - gran_total_costo)

        for c in [c_fin_1, c_fin_2, c_fin_3, c_fin_4]:
            c.font = FONT_TOTAL; c.fill = FILL_TOTAL; c.border = BORDER_TOTAL
        c_fin_2.number_format = "$#,##0"
        c_fin_3.number_format = "$#,##0"
        c_fin_4.number_format = "$#,##0"

        _aplicar_autoajuste_columnas(ws_balance)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output
